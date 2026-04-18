import logging
import os
import subprocess
import tempfile

from app.core.config import settings

logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Shared utilities
# ---------------------------------------------------------------------------


def extract_audio_from_video(video_path: str) -> str:
    """Extract audio track from video file using ffmpeg."""
    audio_path = tempfile.mktemp(suffix=".wav")
    cmd = [
        "ffmpeg", "-i", video_path,
        "-vn", "-acodec", "pcm_s16le",
        "-ar", "16000", "-ac", "1",
        audio_path, "-y",
    ]
    result = subprocess.run(cmd, capture_output=True, text=True)
    if result.returncode != 0:
        raise RuntimeError(f"ffmpeg error: {result.stderr}")
    return audio_path


def read_text_file(file_path: str) -> str:
    """Read text from .txt or .docx file."""
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".docx":
        from docx import Document
        doc = Document(file_path)
        return "\n".join(p.text for p in doc.paragraphs)
    with open(file_path, "r", encoding="utf-8") as f:
        return f.read()


# ---------------------------------------------------------------------------
# Provider: local (openai-whisper + pyannote)
# ---------------------------------------------------------------------------


def _transcribe_local(audio_path: str) -> str:
    """Transcribe with local Whisper + pyannote speaker diarization."""
    import whisper

    model_name = settings.WHISPER_MODEL
    model = whisper.load_model(model_name)
    whisper_result = model.transcribe(audio_path, language=None)
    whisper_segments = whisper_result.get("segments", [])

    try:
        diarization_segments = _diarize_audio(audio_path)
    except Exception:
        logger.warning("Diarization failed, returning plain transcript")
        return whisper_result["text"]

    lines: list[str] = []
    for ws in whisper_segments:
        seg_mid = (ws["start"] + ws["end"]) / 2
        speaker = _find_speaker(diarization_segments, seg_mid)
        lines.append(f"[{speaker}]: {ws['text'].strip()}")

    return "\n".join(lines)


def _diarize_audio(audio_path: str) -> list[dict]:
    from pyannote.audio import Pipeline

    pipeline = Pipeline.from_pretrained("pyannote/speaker-diarization-3.1")
    diarization = pipeline(audio_path)

    segments = []
    for turn, _, speaker in diarization.itertracks(yield_label=True):
        segments.append({
            "speaker": speaker,
            "start": turn.start,
            "end": turn.end,
        })
    return segments


def _find_speaker(segments: list[dict], time_point: float) -> str:
    for seg in segments:
        if seg["start"] <= time_point <= seg["end"]:
            return seg["speaker"]
    return "Неизвестный"


# ---------------------------------------------------------------------------
# Provider: openai (Whisper API)
# ---------------------------------------------------------------------------


def _transcribe_openai(audio_path: str) -> str:
    """Transcribe via OpenAI Whisper API. No diarization."""
    from openai import OpenAI

    client = OpenAI(api_key=settings.OPENAI_API_KEY)
    with open(audio_path, "rb") as f:
        result = client.audio.transcriptions.create(model="whisper-1", file=f)
    return result.text


# ---------------------------------------------------------------------------
# Provider: assemblyai
# ---------------------------------------------------------------------------


def _transcribe_assemblyai(audio_path: str) -> str:
    """Transcribe via AssemblyAI with built-in speaker diarization."""
    import assemblyai as aai

    aai.settings.api_key = settings.ASSEMBLYAI_API_KEY

    config = aai.TranscriptionConfig(speaker_labels=True, language_code="ru")
    transcriber = aai.Transcriber()
    transcript = transcriber.transcribe(audio_path, config=config)

    if transcript.status == aai.TranscriptStatus.error:
        raise RuntimeError(f"AssemblyAI error: {transcript.error}")

    if transcript.utterances:
        lines: list[str] = []
        for u in transcript.utterances:
            lines.append(f"[Speaker {u.speaker}]: {u.text}")
        return "\n".join(lines)

    return transcript.text or ""


# ---------------------------------------------------------------------------
# Dispatch
# ---------------------------------------------------------------------------

_PROVIDERS = {
    "local": _transcribe_local,
    "openai": _transcribe_openai,
    "assemblyai": _transcribe_assemblyai,
}


def transcribe(audio_path: str, provider: str | None = None) -> str:
    """Transcribe audio using the specified (or default) provider."""
    provider = provider or settings.TRANSCRIPTION_PROVIDER
    fn = _PROVIDERS.get(provider)
    if fn is None:
        raise ValueError(f"Unknown transcription provider: {provider}")
    return fn(audio_path)
